from io import BytesIO
import json
from pathlib import Path
import sys
import tempfile
import unittest
import numpy as np
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'app'))
from rag_documents import Section, chunk_sections, extract_document
from rag_store import Store
from rag_answer import evidence_for_prompt, messages_for, validate_answer


class RagTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.store = Store(Path(self.tmp.name) / 'index.sqlite3')

    def tearDown(self):
        self.tmp.cleanup()

    def add(self):
        chunks = [{'location':'PDFページ2','text':'設定確認後に試験を開始する。'},
                  {'location':'PDFページ3','text':'接続切断時は試験を一時停止する。'}]
        return self.store.add_document('abc','仕様書.pdf',chunks,[[1,0],[0,1]],'fake@v1')

    def test_search_returns_correct_document_and_page(self):
        self.add()
        hits = self.store.search([0,1],'fake@v1',1)
        self.assertEqual(hits[0]['location'],'PDFページ3')
        self.assertEqual(hits[0]['name'],'仕様書.pdf')
        self.assertAlmostEqual(hits[0]['score'],1)

    def test_duplicate_is_not_reindexed(self):
        self.assertTrue(self.add())
        self.assertFalse(self.add())
        self.assertEqual(self.store.documents()[0]['chunks'],2)

    def test_changed_embedding_rejected_without_partial_write(self):
        self.add()
        with self.assertRaises(ValueError):
            self.store.add_document('xyz','別文書.txt',[{'location':'行1','text':'別'}],[[1,0]],'other')
        self.assertEqual(len(self.store.documents()),1)
        with self.assertRaises(ValueError):
            self.store.search([1,0],'other')

    def test_document_filter(self):
        self.add()
        self.assertEqual(self.store.search([1,0],'fake@v1',document_ids=['nonexistent']),[])

    def test_invalid_vector_rejected(self):
        for vector in ([0,0],[float('nan'),1]):
            with self.assertRaises(ValueError):
                Store.vectors([vector])

    def test_chunking_preserves_all_characters_and_location(self):
        text = ''.join(chr(0x4e00+i) for i in range(300))
        chunks = chunk_sections([Section('ページ1',text)],100,20,token_count=len,max_tokens=55)
        self.assertTrue(all(len(c['text'])<=55 for c in chunks))
        self.assertEqual(chunks[0]['text'][0],text[0])
        self.assertEqual(chunks[-1]['text'][-1],text[-1])
        self.assertTrue(set(text) <= set(''.join(c['text'] for c in chunks)))
        self.assertTrue(all(c['location']=='ページ1' for c in chunks))

    def test_txt_extract_uses_real_line_locations(self):
        sections,warnings = extract_document('notes.txt','\n'.join(f'行{i}' for i in range(25)).encode())
        self.assertEqual([s.location for s in sections],['行1-20','行21-25'])
        self.assertEqual(warnings,[])
        with self.assertRaises(ValueError):
            extract_document('a.exe',b'abcd')

    def test_docx_paragraph_and_table(self):
        from docx import Document
        document=Document()
        document.add_paragraph('接続の確認')
        table=document.add_table(rows=1,cols=2)
        table.cell(0,0).text='項目';table.cell(0,1).text='値'
        raw=BytesIO();document.save(raw)
        sections,warnings=extract_document('doc.docx',raw.getvalue())
        self.assertEqual(sections[0].location,'本文段落1')
        self.assertEqual(sections[1].text,'項目 | 値')
        self.assertTrue(warnings)

    def test_pdf_text_and_scanned_page_warning(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        writer = PdfWriter()
        page = writer.add_blank_page(612, 792)
        font = DictionaryObject({NameObject('/Type'):NameObject('/Font'),
                                 NameObject('/Subtype'):NameObject('/Type1'),
                                 NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):
            DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
        content = DecodedStreamObject()
        content.set_data(b'BT /F1 12 Tf 50 700 Td (Document text) Tj ET')
        page[NameObject('/Contents')] = writer._add_object(content)
        writer.add_blank_page(612,792)
        raw=BytesIO();writer.write(raw)
        sections,warnings=extract_document('doc.pdf',raw.getvalue())
        self.assertEqual(sections[0].location,'PDFページ1')
        self.assertTrue(any('ページ2' in w for w in warnings))

    def test_prompt_budget_and_low_similarity(self):
        hits=[{'score':.9,'name':'仕様書','location':'p1','text':'abc'},
              {'score':.3,'name':'仕様書','location':'p2','text':'xyz'}]
        evidence=evidence_for_prompt(hits,.75,lambda _:100,100)
        self.assertEqual(len(evidence),1)
        self.assertEqual(evidence[0]['citation'],1)
        self.assertEqual(evidence_for_prompt(hits,.75,lambda _:101,100),[])
        messages=messages_for('内容を説明',evidence)
        self.assertIn('信頼できない入力',messages[0]['content'])
        self.assertEqual(json.loads(messages[1]['content'])['reference_data'][0]['location'],'p1')

    def test_citation_validation(self):
        refs=[{'citation':1}]
        self.assertIsNone(validate_answer('確認します。[1]',refs))
        self.assertIsNotNone(validate_answer('確認します。[2]',refs))
        self.assertIsNotNone(validate_answer('確認します。',refs))
        self.assertIsNone(validate_answer('資料では確認できません',refs))

    def test_review_requires_evidence_and_exports_grounded_example(self):
        sources=[{'document_id':'abc','name':'仕様書','location':'p1','text':'設定を確認する','citation':1}]
        self.store.save_review('条件は？','設定を確認します。[1]',sources)
        self.store.save_review('条件は？','設定を確認します。[1]',sources)
        self.assertEqual(len(self.store.reviews()),1)
        row=json.loads(self.store.export_reviews())
        self.assertIn('設定を確認する',row['messages'][1]['content'])
        self.assertEqual(row['source_metadata'],sources)
        self.assertTrue(row['reviewed'])
        with self.assertRaises(ValueError):
            self.store.save_review('q','a',[])

    def test_training_rejects_same_source_across_splits(self):
        from data_utils import validate_directory
        sources=[{'document_id':'same-document','name':'仕様書','location':'p1','text':'原文'}]
        for split in ('train','valid','test'):
            row={'messages':[{'role':'user','content':split+'の質問'},
                             {'role':'assistant','content':'回答'}], 'source_metadata':sources}
            Path(self.tmp.name,split+'.jsonl').write_text(json.dumps(row)+'\n')
        with self.assertRaisesRegex(ValueError,'同じ元文書'):
            validate_directory(self.tmp.name)


if __name__ == '__main__':
    unittest.main()
