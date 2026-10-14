"""Local document extraction; source locations never invent Word page numbers."""
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
import zipfile

MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_TEXT_CHARS = 2_000_000


@dataclass(frozen=True)
class Section:
    location: str
    text: str


def extract_document(name, raw):
    if not raw or len(raw) > MAX_FILE_BYTES:
        raise ValueError('空ファイル、または50MBを超えるファイルは登録できません')
    ext = Path(name).suffix.lower()
    sections, warnings = [], []
    if ext in ('.txt', '.md'):
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError as exc:
            raise ValueError('TXT/MDはUTF-8で保存してください') from exc
        lines = text.splitlines()
        # Bounded line groups retain explicit source positions.
        for start in range(0, len(lines), 20):
            sections.append(Section(f'行{start+1}-{min(start+20, len(lines))}', '\n'.join(lines[start:start+20])))
    elif ext == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(raw))
        if reader.is_encrypted:
            raise ValueError('暗号化PDFは未対応です。許可された復号済み文書を使用してください')
        if len(reader.pages) > 500:
            raise ValueError('PDFは500ページまでです。章単位に分割してください')
        count = 0
        for i, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ''
            count += len(text)
            if count > MAX_TEXT_CHARS:
                raise ValueError('抽出テキストが200万文字を超えます。文書を分割してください')
            if not text.strip():
                warnings.append(f'PDFページ{i}: テキストがありません。画像・スキャン部分はOCRが必要です')
            else:
                sections.append(Section(f'PDFページ{i}', text))
    elif ext == '.docx':
        from docx import Document
        # Guard against decompression bombs before XML parsing.
        with zipfile.ZipFile(BytesIO(raw)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('DOCXの展開サイズが大きすぎます')
        document = Document(BytesIO(raw))
        for i, para in enumerate(document.paragraphs, 1):
            if para.text.strip():
                sections.append(Section(f'本文段落{i}', para.text))
        for i, table in enumerate(document.tables, 1):
            for j, row in enumerate(table.rows, 1):
                text = ' | '.join(cell.text for cell in row.cells)
                if text.strip(' |'):
                    sections.append(Section(f'表{i}・行{j}', text))
        warnings.append('Wordのページ番号・画像・ヘッダー・脚注・テキストボックスは未取得です。本文段落と表の位置で引用します')
    else:
        raise ValueError('対応形式はPDF / DOCX / TXT / MDです')
    sections = [s for s in sections if s.text.strip()]
    if sum(len(s.text) for s in sections) > MAX_TEXT_CHARS:
        raise ValueError('抽出テキストが200万文字を超えます。文書を分割してください')
    if not sections:
        raise ValueError('本文を抽出できません。スキャン文書は先にOCRしてください')
    return sections, warnings


def chunk_sections(sections, max_chars=600, overlap=80, token_count=None, max_tokens=450):
    if max_chars < 32 or not 0 <= overlap < max_chars:
        raise ValueError('チャンク長と重複長の設定が不正です')
    chunks = []
    for section in sections:
        text = section.text.strip()
        start = 0
        while start < len(text):
            end = min(start + max_chars, len(text))
            while token_count is not None and token_count(text[start:end]) > max_tokens:
                end = start + (end - start) // 2
                if end <= start:
                    raise ValueError('テキストを埋め込みトークン上限内に分割できません')
            excerpt = text[start:end].strip()
            if excerpt:
                chunks.append({'location': section.location, 'text': excerpt})
            if end == len(text):
                break
            start = max(start + 1, end - min(overlap, (end - start) // 4))
    return chunks
